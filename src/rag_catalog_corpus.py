"""
ALCheMy RAG - Corpus Cataloguer (Phase A-bis)

Walks data/rag_corpus/{public,internal}, extracts bibliographic metadata from
every PDF/DOCX, and fills the `rag_documents` table. It does NOT create
embeddings: vectorisation happens in the n8n workflow. This script is the
*card catalogue*; n8n is the *shelf*.

Why it exists
-------------
The binary MD5 of a file is a perfect fingerprint of the BYTES but a terrible
fingerprint of the PAPER: publishers stamp a download watermark (date, user)
on every page, so the same article downloaded twice yields two different
hashes. That watermark is real text inside the PDF, so hashing the extracted
text does not save you either - both were verified to break.

Hence a three-level ladder, implemented in SQL as rag_find_existing_document():

    level 1  binary MD5      -> the literally identical file
    level 2  DOI             -> the same paper from another source/watermark
    level 3  header text MD5 -> documents with no DOI (page proofs, notebooks)

Usage
-----
    python src/rag_catalog_corpus.py                # dry run: report only
    python src/rag_catalog_corpus.py --apply        # write into rag_documents
    python src/rag_catalog_corpus.py --apply --crossref
        Additionally resolves each DOI against the public Crossref API to fill
        csl_json. Only PUBLIC documents are ever sent, and only their DOI - an
        already-published identifier. INTERNAL/RESTRICTED material never leaves
        the machine, per the confidentiality rules in CLAUDE.md.
"""

import argparse
import hashlib
import io
import json
import os
import re
import shutil
import sys
import unicodedata
from pathlib import Path

import fitz  # PyMuPDF
import psycopg2
import psycopg2.extras

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #
REPO_ROOT = Path(__file__).resolve().parent.parent
CORPUS_ROOT = REPO_ROOT / "data" / "rag_corpus"

# The sub-folder decides the confidentiality tier. Nobody has to remember to
# tick a box: dropping a file in internal/ is the act of labelling it.
SENSITIVITY_BY_FOLDER = {"public": "PUBLIC", "internal": "INTERNAL",
                         "restricted": "RESTRICTED"}

SUPPORTED_EXT = {".pdf", ".docx", ".txt", ".md"}

# A DOI always starts with "10.", a registrant code, a slash, and a suffix.
DOI_RE = re.compile(r"\b10\.\d{4,9}/[-._;()/:A-Za-z0-9]+\b")

# Filename conventions that mark a Supporting Information file.
SI_RE = re.compile(r"(_si\b|\bsi_|supporting|supp[-_ ]?info|electronic[-_ ]?supp)",
                   re.IGNORECASE)

HEADER_CHARS = 3000  # how much of the document feeds the level-3 hash

# Below this many extractable characters a document is treated as a scan: it
# would be indexed without error and yield zero chunks, which is the worst
# kind of failure because nothing looks broken.
MIN_CHARS_FOR_TEXT = 200


# --------------------------------------------------------------------------- #
# Text extraction
# --------------------------------------------------------------------------- #
def normalise(text):
    """Collapse whitespace, unify unicode forms, lowercase.

    This is what makes the level-3 hash survive cosmetic differences: a PDF
    re-exported by another tool may change line breaks and ligatures without
    changing a single word.
    """
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text)).strip().lower()


def read_pdf(path):
    doc = fitz.open(path)
    pages = [page.get_text() for page in doc]
    n_images = sum(len(doc[i].get_images(full=True)) for i in range(doc.page_count))
    meta = doc.metadata or {}
    result = {
        "text": "".join(pages),
        "n_pages": doc.page_count,
        "n_images": n_images,
        "meta_title": (meta.get("title") or "").strip(),
        "meta_author": (meta.get("author") or "").strip(),
    }
    doc.close()
    return result


def read_docx(path):
    from docx import Document  # imported lazily: only Word files need it

    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs]
    # Tables in a Supporting Information often hold the synthesis conditions,
    # so they must not be dropped.
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    core = doc.core_properties
    return {
        "text": "\n".join(paragraphs),
        "n_pages": None,  # .docx has no fixed pagination
        "n_images": len(doc.inline_shapes),
        "meta_title": (core.title or "").strip(),
        "meta_author": (core.author or "").strip(),
    }


def read_plain(path):
    text = path.read_text(encoding="utf-8", errors="replace")
    return {"text": text, "n_pages": None, "n_images": 0,
            "meta_title": "", "meta_author": ""}


READERS = {".pdf": read_pdf, ".docx": read_docx, ".txt": read_plain, ".md": read_plain}


# --------------------------------------------------------------------------- #
# Metadata inference
# --------------------------------------------------------------------------- #
def find_doi(text):
    """Search the first pages only: the reference list is full of OTHER papers'
    DOIs, and picking one of those would mislabel the document."""
    for match in DOI_RE.findall(text[:6000]):
        return match.rstrip(".,;)")
    return None


def guess_year(text):
    """A 4-digit year between 1900 and 2100 near the head of the document."""
    for candidate in re.findall(r"\b(19\d{2}|20\d{2})\b", text[:6000]):
        year = int(candidate)
        if 1900 <= year <= 2100:
            return year
    return None


def guess_title(meta_title, text, file_name):
    """PDF metadata first - when the publisher filled it in, it is exact.
    Otherwise the longest of the first few non-trivial lines, which in a
    typeset paper is almost always the title."""
    if meta_title and len(meta_title) > 10:
        return meta_title
    lines = [line.strip() for line in text[:2000].splitlines() if len(line.strip()) > 15]
    return max(lines[:12], key=len) if lines else Path(file_name).stem


def file_md5(path):
    digest = hashlib.md5()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def extract(path, sensitivity):
    ext = path.suffix.lower()
    raw = READERS[ext](path)
    text = raw["text"]
    doi = find_doi(text)
    title = guess_title(raw["meta_title"], text, path.name)

    # Level-3 fingerprint: title + the head of the document, normalised. Using
    # only the head means a watermark stamped on page 7 no longer breaks it.
    header = normalise(title + " " + text[:HEADER_CHARS])
    # Only the file name and the TITLE decide. Scanning the body would flag
    # every paper as its own Supporting Information, since papers routinely
    # say "see the Supporting Information" in their abstract and footnotes.
    is_si = bool(SI_RE.search(path.stem)) or bool(SI_RE.search(title[:120]))

    return {
        "file_name": path.name,
        # Container path, not host path: n8n and the runners see it this way.
        "file_path": "/home/node/data/rag_corpus/{}/{}".format(
            sensitivity.lower(), path.name),
        "file_ext": ext.lstrip("."),
        "file_hash": file_md5(path),
        "text_hash": hashlib.md5(header.encode("utf-8")).hexdigest(),
        "doi": doi,
        "title": title,
        "authors": raw["meta_author"] or None,
        "year": guess_year(text),
        "journal": None,
        "doc_type": ("SUPPORTING_INFORMATION" if is_si
                     else "LAB_NOTEBOOK" if sensitivity != "PUBLIC"
                     else "PAPER"),
        "sensitivity": sensitivity,
        "n_pages": raw["n_pages"],
        "n_images": raw["n_images"],
        # Everything guessed from the file itself is the weakest tier.
        "metadata_source": "PDF_HEURISTIC",
        "n_chars": len(text.strip()),
        "needs_ocr": len(text.strip()) < MIN_CHARS_FOR_TEXT,
        "csl_json": {},
        "_is_si": is_si,
    }


# --------------------------------------------------------------------------- #
# Crossref enrichment (PUBLIC documents only)
# --------------------------------------------------------------------------- #
def enrich_from_crossref(record):
    """Resolve a DOI into CSL-JSON. Sends nothing but the DOI, which is public
    by definition. Guarded by the caller so it never runs on INTERNAL files."""
    import requests

    url = "https://api.crossref.org/works/{}".format(record["doi"])
    try:
        response = requests.get(
            url, timeout=15,
            headers={"User-Agent": "ALCheMy-RAG/1.0 (mailto:rubbenserrano@gmail.com)"},
        )
        if response.status_code != 200:
            return False
        work = response.json().get("message", {})
    except Exception:                                   # noqa: BLE001
        return False

    authors = ", ".join(
        "{} {}".format(a.get("given", ""), a.get("family", "")).strip()
        for a in work.get("author", [])
    )
    parts = (work.get("issued", {}).get("date-parts") or [[None]])[0]

    if work.get("title"):
        record["title"] = work["title"][0]
    if authors:
        record["authors"] = authors
    if work.get("container-title"):
        record["journal"] = work["container-title"][0]
    if parts and parts[0]:
        record["year"] = int(parts[0])
    record["csl_json"] = work
    record["metadata_source"] = "CROSSREF"
    return True


# --------------------------------------------------------------------------- #
# Database
# --------------------------------------------------------------------------- #
def connect():
    try:
        from dotenv import load_dotenv       # optional dependency
        load_dotenv(REPO_ROOT / ".env")
    except Exception:                                   # noqa: BLE001
        _load_env_manually(REPO_ROOT / ".env")
    return psycopg2.connect(
        host=os.getenv("POSTGRES_HOST", "localhost"),
        port=os.getenv("POSTGRES_PORT", "5432"),
        dbname=os.environ["POSTGRES_DB"],
        user=os.environ["POSTGRES_USER"],
        password=os.environ["POSTGRES_PASSWORD"],
    )


def _load_env_manually(path):
    """Minimal .env reader so the script runs without python-dotenv."""
    if not path.is_file():
        return
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def already_indexed(cursor, record):
    cursor.execute(
        "SELECT document_id, file_name, motivo "
        "FROM rag_find_existing_document(%s, %s, %s);",
        (record["file_hash"], record["doi"], record["text_hash"]),
    )
    return cursor.fetchone()


UPSERT = """
INSERT INTO rag_documents
    (file_name, file_path, file_ext, file_hash, text_hash, doi, title, authors,
     year, journal, doc_type, sensitivity, n_pages, n_images, n_chars,
     needs_ocr, metadata_source, csl_json)
VALUES
    (%(file_name)s, %(file_path)s, %(file_ext)s, %(file_hash)s, %(text_hash)s,
     %(doi)s, %(title)s, %(authors)s, %(year)s, %(journal)s, %(doc_type)s,
     %(sensitivity)s, %(n_pages)s, %(n_images)s, %(n_chars)s,
     %(needs_ocr)s, %(metadata_source)s, %(csl_json)s)
ON CONFLICT (file_hash) DO UPDATE SET
    -- file_name and file_path are deliberately NOT refreshed. Two different
    -- Zotero entries can share the exact same extracted text, hence the same
    -- file_hash; the second one must not repoint the row at a file that was
    -- never written to disk precisely because it was a duplicate. The path
    -- recorded the first time is the one that exists.
    -- Physical measurements are refreshed: they are facts, not opinions, and
    -- re-reading the same bytes cannot make them worse.
    file_ext  = EXCLUDED.file_ext,
    text_hash = EXCLUDED.text_hash,
    n_pages   = EXCLUDED.n_pages,
    n_images  = EXCLUDED.n_images,
    n_chars   = EXCLUDED.n_chars,
    needs_ocr = EXCLUDED.needs_ocr,
    -- Bibliographic fields are opinions, and they only move upwards: a weaker
    -- source must never overwrite a stronger one. This is what stopped the
    -- single-author PDF metadata from clobbering the nine authors of Crossref.
    doi     = CASE WHEN rag_source_rank(EXCLUDED.metadata_source)
                      >= rag_source_rank(rag_documents.metadata_source)
                   THEN COALESCE(EXCLUDED.doi, rag_documents.doi)
                   ELSE rag_documents.doi END,
    title   = CASE WHEN rag_source_rank(EXCLUDED.metadata_source)
                      >= rag_source_rank(rag_documents.metadata_source)
                   THEN COALESCE(EXCLUDED.title, rag_documents.title)
                   ELSE rag_documents.title END,
    authors = CASE WHEN rag_source_rank(EXCLUDED.metadata_source)
                      >= rag_source_rank(rag_documents.metadata_source)
                   THEN COALESCE(EXCLUDED.authors, rag_documents.authors)
                   ELSE rag_documents.authors END,
    year    = CASE WHEN rag_source_rank(EXCLUDED.metadata_source)
                      >= rag_source_rank(rag_documents.metadata_source)
                   THEN COALESCE(EXCLUDED.year, rag_documents.year)
                   ELSE rag_documents.year END,
    journal = CASE WHEN rag_source_rank(EXCLUDED.metadata_source)
                      >= rag_source_rank(rag_documents.metadata_source)
                   THEN COALESCE(EXCLUDED.journal, rag_documents.journal)
                   ELSE rag_documents.journal END,
    csl_json = CASE WHEN EXCLUDED.csl_json <> '{}'::jsonb
                     AND rag_source_rank(EXCLUDED.metadata_source)
                      >= rag_source_rank(rag_documents.metadata_source)
                    THEN EXCLUDED.csl_json ELSE rag_documents.csl_json END,
    metadata_source = CASE WHEN rag_source_rank(EXCLUDED.metadata_source)
                              >= rag_source_rank(rag_documents.metadata_source)
                           THEN EXCLUDED.metadata_source
                           ELSE rag_documents.metadata_source END
RETURNING document_id;
"""

# A Supporting Information carries the DOI of the article it belongs to, so the
# link is drawn on that. The WHERE clause excludes self-parenting and only ever
# hangs an SI off a PAPER.
LINK_SI = """
UPDATE rag_documents si
SET    parent_document_id = paper.document_id
FROM   rag_documents paper
WHERE  si.doc_type = 'SUPPORTING_INFORMATION'
  AND  paper.doc_type = 'PAPER'
  AND  si.doi IS NOT NULL
  AND  lower(si.doi) = lower(paper.doi)
  AND  si.document_id <> paper.document_id
  AND  si.parent_document_id IS DISTINCT FROM paper.document_id;
"""


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
UPGRADE = """
UPDATE rag_documents SET
    doi     = CASE WHEN rag_source_rank(%(metadata_source)s)
                      >= rag_source_rank(metadata_source)
                   THEN COALESCE(%(doi)s, doi) ELSE doi END,
    title   = CASE WHEN rag_source_rank(%(metadata_source)s)
                      >= rag_source_rank(metadata_source)
                   THEN COALESCE(%(title)s, title) ELSE title END,
    authors = CASE WHEN rag_source_rank(%(metadata_source)s)
                      >= rag_source_rank(metadata_source)
                   THEN COALESCE(%(authors)s, authors) ELSE authors END,
    year    = CASE WHEN rag_source_rank(%(metadata_source)s)
                      >= rag_source_rank(metadata_source)
                   THEN COALESCE(%(year)s, year) ELSE year END,
    journal = CASE WHEN rag_source_rank(%(metadata_source)s)
                      >= rag_source_rank(metadata_source)
                   THEN COALESCE(%(journal)s, journal) ELSE journal END,
    csl_json = CASE WHEN %(csl_json)s::jsonb <> '{}'::jsonb
                     AND rag_source_rank(%(metadata_source)s)
                      >= rag_source_rank(metadata_source)
                    THEN %(csl_json)s::jsonb ELSE csl_json END,
    metadata_source = CASE WHEN rag_source_rank(%(metadata_source)s)
                              >= rag_source_rank(metadata_source)
                           THEN %(metadata_source)s ELSE metadata_source END
WHERE document_id = %(document_id)s;
"""


def _upgrade(cursor, document_id, record):
    """The same paper arriving as a DIFFERENT file (levels 2 and 3 of the
    ladder). Inserting would duplicate the work, and the unique DOI index would
    reject it anyway. What is worth doing is lifting the existing card with the
    better metadata - Zotero outranks Crossref and the PDF heuristics.
    """
    payload = {k: record.get(k) for k in
               ("doi", "title", "authors", "year", "journal", "metadata_source")}
    payload["csl_json"] = json.dumps(record.get("csl_json") or {})
    payload["document_id"] = document_id
    cursor.execute(UPGRADE, payload)


def _write(cursor, record):
    """UPSERT one card. Keyed on file_hash, so calling it twice is harmless."""
    payload = dict(record)
    payload.pop("_is_si", None)
    payload.pop("_text", None)
    payload["csl_json"] = json.dumps(payload["csl_json"])
    cursor.execute(UPSERT, payload)


# Zotero collections whose content is not chemistry. Excluded by default; the
# owner's library also holds a dental systematic review and a fitness folder.
EXCLUDED_COLLECTIONS_DEFAULT = (
    "CBCT para prevencion en sinfisis mandibular",
    "FITNESS",
    "26_PRISMA_Incluidos_Final_35",
    "32567942",
    "32567942 (1)",
)


def _corpus_name(item, extension):
    """A readable, collision-resistant file name: FirstAuthor_Year_key.ext

    The Zotero key is what makes it unique - two papers by the same author in
    the same year are common, and storage folder names are opaque.
    """
    first = (item["csl"]["author"] or [{}])[0].get("family") or "anon"
    first = re.sub(r"[^A-Za-z0-9]+", "", first)[:24] or "anon"
    return "{}_{}_{}{}".format(first, item["year"] or "sf", item["zotero_key"], extension)


def _record_from_text(item, text, target, subfolder):
    """Build a catalogue record straight from Zotero's extracted text.

    No PDF is involved, so page and image counts are unknown - they are
    properties of a file we do not have. Everything that matters for retrieval
    (the text) and for citing (the metadata) is present.
    """
    normalised = normalise((item["title"] or "") + " " + text[:HEADER_CHARS])
    return {
        "file_name": target.name,
        "file_path": "/home/node/data/rag_corpus/{}/{}".format(subfolder, target.name),
        "file_ext": "txt",
        "file_hash": hashlib.md5(text.encode("utf-8", "replace")).hexdigest(),
        "text_hash": hashlib.md5(normalised.encode("utf-8")).hexdigest(),
        "doi": item["doi"],
        "title": item["title"],
        "authors": item["authors"],
        "year": item["year"],
        "journal": item["journal"],
        "doc_type": "PAPER",
        "sensitivity": "PUBLIC",
        "n_pages": None,
        "n_images": None,
        "n_chars": len(text.strip()),
        "needs_ocr": len(text.strip()) < MIN_CHARS_FOR_TEXT,
        "metadata_source": "REFERENCE_MANAGER",
        "csl_json": item["csl"],
        "_is_si": False,
        "_text": text,
    }


def zotero_records(zotero_dir, subfolder, excluded, use_text_cache):
    """Second ingestion lane: a reference manager instead of a folder.

    Dropping PDFs in a folder is convenient but leans on heuristics, and
    publishers fill PDF metadata badly. Zotero already holds the record curated
    by a human, so these entries arrive as REFERENCE_MANAGER and outrank both
    the heuristics and Crossref.

    Two sources of text, in order of preference:
      1. the PDF, when it is still on disk (also gives page and image counts);
      2. Zotero's .zotero-ft-cache, the text it extracted when the PDF *was*
         on disk. Most of this library synced its metadata but not its files,
         so without this second source almost nothing would be indexable.
    """
    import rag_zotero

    for item in rag_zotero.read_library(zotero_dir, require_pdf=False,
                                        require_text=True,
                                        excluded_collections=excluded):
        destination = CORPUS_ROOT / subfolder
        if item["pdf_exists"]:
            source = Path(item["pdf_path"])
            target = destination / _corpus_name(item, ".pdf")
            record = extract(source, "PUBLIC")
            record["file_name"] = target.name
            record["file_path"] = "/home/node/data/rag_corpus/{}/{}".format(
                subfolder, target.name)
            for field in ("title", "authors", "year", "journal", "doi"):
                if item.get(field):
                    record[field] = item[field]
            record["csl_json"] = item["csl"]
            record["metadata_source"] = "REFERENCE_MANAGER"
            record["_text"] = None
            yield record, source, target
        elif use_text_cache and item["text_cache"]:
            target = destination / _corpus_name(item, ".txt")
            with io.open(item["text_cache"], encoding="utf-8", errors="replace") as handle:
                text = handle.read()
            yield _record_from_text(item, text, target, subfolder), None, target


def collect_files():
    for folder, sensitivity in SENSITIVITY_BY_FOLDER.items():
        directory = CORPUS_ROOT / folder
        if not directory.is_dir():
            continue
        for path in sorted(directory.iterdir()):
            if path.is_file() and path.suffix.lower() in SUPPORTED_EXT:
                yield path, sensitivity


def main():
    parser = argparse.ArgumentParser(description="ALCheMy RAG corpus cataloguer")
    parser.add_argument("--apply", action="store_true",
                        help="write to rag_documents (default is a dry run)")
    parser.add_argument("--crossref", action="store_true",
                        help="resolve DOIs of PUBLIC documents via Crossref")
    parser.add_argument("--zotero", metavar="DIR",
                        help="also read a Zotero library (the folder holding "
                             "zotero.sqlite and storage/)")
    parser.add_argument("--copy-into", metavar="SUBFOLDER", default="zotero",
                        help="with --zotero and --apply, write each document into "
                             "data/rag_corpus/<SUBFOLDER> (default: zotero)")
    parser.add_argument("--no-text-cache", action="store_true",
                        help="with --zotero, use only entries whose PDF is on disk")
    parser.add_argument("--exclude-collections", metavar="A,B",
                        default=",".join(EXCLUDED_COLLECTIONS_DEFAULT),
                        help="top-level Zotero collections to skip (comma separated)")
    args = parser.parse_args()

    files = list(collect_files())
    if not files and not args.zotero:
        print("No documents found under {}".format(CORPUS_ROOT))
        return 0

    connection = connect()
    connection.autocommit = False
    cursor = connection.cursor(cursor_factory=psycopg2.extras.DictCursor)

    new, duplicates, written = [], [], 0

    for path, sensitivity in files:
        try:
            record = extract(path, sensitivity)
        except Exception as error:                      # noqa: BLE001
            print("  !! {:<42s} unreadable: {}".format(path.name[:42], error))
            continue

        # Enrich BEFORE the duplicate check, so that a file already on the
        # shelf also gets its card upgraded. The UPSERT ranks sources, so this
        # can only improve the record, never degrade it.
        if args.crossref and record["doi"] and sensitivity == "PUBLIC":
            if enrich_from_crossref(record):
                print("      -> Crossref: {} | {} | {} autores".format(
                    (record["journal"] or "?")[:40], record["year"],
                    len(record["csl_json"].get("author") or [])))

        hit = already_indexed(cursor, record)
        flag = "SI " if record["_is_si"] else "   "
        print("{} {:<42s} {:<9s} {}  pags={:<5} imgs={:<5} DOI={}".format(
            flag, path.name[:42], sensitivity,
            "DUP" if hit else "NEW",
            record["n_pages"] if record["n_pages"] is not None else "-",
            record["n_images"],
            record["doi"] or "-"))

        if record["needs_ocr"]:
            print("      !! SIN TEXTO EXTRAIBLE ({} caracteres): es un escaneo. "
                  "Marcado needs_ocr=TRUE; la ingesta debe saltarselo hasta "
                  "transcribirlo con un VLM.".format(record["n_chars"]))

        if hit:
            print("      -> ya indexado como '{}' ({})".format(
                hit["file_name"], hit["motivo"]))
            duplicates.append(record)
            # Level 1 means it is byte-for-byte the same file, so re-reading it
            # can only refresh the card with better metadata - the UPSERT is
            # keyed on file_hash and cannot create a second row. Levels 2 and 3
            # mean a DIFFERENT file carrying the same paper: that one must not
            # be written, or the corpus would hold the article twice.
            if args.apply:
                if hit["motivo"].startswith("nivel 1"):
                    _write(cursor, record)      # mismo fichero: refrescar
                else:
                    _upgrade(cursor, hit["document_id"], record)
            continue
        new.append(record)

        if args.apply:
            _write(cursor, record)
            written += 1

    # ---- lane 2: Zotero --------------------------------------------------
    if args.zotero:
        excluded = {c.strip() for c in args.exclude_collections.split(",") if c.strip()}
        print("\n--- Biblioteca Zotero: {} ---".format(args.zotero))
        print("    colecciones excluidas: {}".format(", ".join(sorted(excluded)) or "ninguna"))
        destination = CORPUS_ROOT / args.copy_into
        if args.apply:
            destination.mkdir(parents=True, exist_ok=True)

        zot_new = zot_dup = zot_vacios = 0
        for record, source, target in zotero_records(
                args.zotero, args.copy_into, excluded, not args.no_text_cache):
            if record["needs_ocr"]:
                zot_vacios += 1
                continue
            hit = already_indexed(cursor, record)
            if hit:
                zot_dup += 1
                if args.apply:
                    if hit["motivo"].startswith("nivel 1"):
                        _write(cursor, record)
                    else:
                        _upgrade(cursor, hit["document_id"], record)
                continue
            zot_new += 1
            if args.apply:
                if record.get("_text") is not None:
                    if not target.exists():
                        with io.open(target, "w", encoding="utf-8") as handle:
                            handle.write(record["_text"])
                elif source is not None and not target.exists():
                    shutil.copy2(str(source), str(target))
                record.pop("_text", None)
                _write(cursor, record)
                written += 1
            if zot_new % 100 == 0:
                print("    ... {} nuevos".format(zot_new))
                if args.apply:
                    connection.commit()
        print("Zotero -> nuevos: {} | ya indexados: {} | sin texto util: {}".format(
            zot_new, zot_dup, zot_vacios))

    if args.apply:
        cursor.execute(LINK_SI)
        linked = cursor.rowcount
        connection.commit()
        print("\nEscritos {} documentos. Supporting Information enlazados: {}."
              .format(written, linked))
    else:
        connection.rollback()
        print("\nDRY RUN - no se ha escrito nada. Repite con --apply.")

    print("Nuevos: {} | Ya indexados: {} | Total inspeccionados: {}".format(
        len(new), len(duplicates), len(files)))

    cursor.close()
    connection.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
